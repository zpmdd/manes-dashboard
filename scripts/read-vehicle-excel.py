"""只读、逐行提取 Excel；stdout 为 NDJSON，不改写原工作簿。"""
import argparse
import datetime
import json
import re
import sys

import openpyxl

FIELDS = dict(zip(
    ['号牌号码', '号牌颜色', '经度', '纬度', '角度', '海拔', 'GPS速度', '行车定位速度', '里程', 'GPS时间', '入库时间', '报警信息', '状态信息', '加密标识', '运营商编码'],
    ['VEHICLENO', 'PALTE_COLOR', 'GEO_LON', 'GEO_LAT', 'GEO_ANG', 'GEO_ALT', 'GPS_SPEED', 'RECORD_SPEED', 'MILEAGE', 'GPS_DATE', 'SI_DT', 'ALARM_CODE', 'STATE_CODE', 'ENCRYPT', 'OPERATOR_CODE']))
REQUIRED = {'VEHICLENO', 'PALTE_COLOR', 'GEO_LON', 'GEO_LAT', 'GPS_SPEED', 'GPS_DATE'}
TEXT_FIELDS = {'VEHICLENO', 'PALTE_COLOR', 'ALARM_CODE', 'STATE_CODE', 'ENCRYPT', 'OPERATOR_CODE'}


def find_header(sheet):
    for number, cells in enumerate(sheet.iter_rows(max_row=20), 1):
        headers = [FIELDS.get(str(c.value).strip(), str(c.value).strip().upper()) if c.value is not None else '' for c in cells]
        if REQUIRED.issubset(headers):
            selected = [h for h in headers if h in FIELDS.values()]
            if len(set(selected)) != len(selected):
                raise ValueError(f'{sheet.title}: 字段重复')
            return number, headers
    return None


def cell_value(cell, key):
    value = cell.value
    if value is None:
        return None
    if cell.data_type in ('f', 'e'):
        raise ValueError(f'{cell.coordinate}: 不接受公式或错误单元格，请提供原始值')
    if isinstance(value, (datetime.datetime, datetime.date)):
        if not isinstance(value, datetime.datetime):
            value = datetime.datetime.combine(value, datetime.time())
        return value.isoformat(sep=' ', timespec='milliseconds')
    if key in TEXT_FIELDS:
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            if key in ('ALARM_CODE', 'STATE_CODE'):
                raise ValueError(f'{cell.coordinate}: 报警/状态码必须存为文本，不能恢复已丢失的前导零')
            if not float(value).is_integer() or abs(value) >= 10**15:
                raise ValueError(f'{cell.coordinate}: 数字标识可能丢失精度，请改用文本')
            text = str(int(value))
            return text.zfill(len(cell.number_format)) if re.fullmatch(r'0+', cell.number_format or '') else text
        return str(value).strip()
    return value


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('file')
    parser.add_argument('--sheet')
    args = parser.parse_args()
    if not args.file.lower().endswith('.xlsx'):
        parser.error('当前接收 .xlsx；旧 .xls 请在 Excel 中另存为 .xlsx')
    workbook = openpyxl.load_workbook(args.file, read_only=True, data_only=False, keep_links=False)
    try:
        candidates = [(s, find_header(s)) for s in workbook if not args.sheet or s.title == args.sheet]
        candidates = [(s, header) for s, header in candidates if header]
        if len(candidates) != 1:
            raise ValueError('需唯一一个包含车辆字段的数据表；多表请用 --sheet 指定，缺表请核对表头')
        sheet, (header_row, headers) = candidates[0]
        for number, cells in enumerate(sheet.iter_rows(min_row=header_row + 1), header_row + 1):
            if all(c.value is None for c in cells):
                continue
            row = {'SOURCE_ROW': f'{sheet.title}!{number}'}
            try:
                for key, cell in zip(headers, cells):
                    if key in FIELDS.values():
                        row[key] = cell_value(cell, key)
            except ValueError as error:
                row['_IMPORT_ERROR'] = str(error)
            print(json.dumps(row, ensure_ascii=False, allow_nan=False))
    finally:
        workbook.close()


if __name__ == '__main__':
    try:
        main()
    except Exception as error:
        print(f'Excel 读取失败：{error}', file=sys.stderr)
        sys.exit(1)
